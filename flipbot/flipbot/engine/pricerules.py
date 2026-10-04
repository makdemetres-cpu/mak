"""Buying-price spreadsheet import (.xlsx / .csv / Google Sheets CSV export).

Prices are never hard-coded. Each import creates a new PriceRuleVersion and
deactivates the previous one, so history keeps which rules applied when.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import PriceRule, PriceRuleVersion
from .parsing import parse_model, parse_storage

FIELDS = ["model", "storage", "max_buy_price", "min_battery_health", "condition_rules", "notes"]
_SYNONYMS = {
    "model": ["model", "μοντελο", "μοντέλο", "device", "phone", "συσκευη"],
    "storage": ["storage", "gb", "χωρητικοτητα", "χωρητικότητα", "capacity", "μνημη"],
    "max_buy_price": ["max buy", "max_buy", "max price", "buy price", "τιμη αγορασ", "τιμή αγοράς",
                      "max", "μεγιστη", "αγορα"],
    "min_battery_health": ["battery", "min battery", "μπαταρια", "μπαταρία", "bh"],
    "condition_rules": ["condition", "κατασταση", "κατάσταση"],
    "notes": ["notes", "σημειωσεισ", "σημειώσεις", "comments"],
}


@dataclass
class ImportReport:
    imported: int = 0
    unmatched_models: list[str] = field(default_factory=list)
    duplicates: list[str] = field(default_factory=list)
    missing_storage: list[str] = field(default_factory=list)
    bad_prices: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return self.__dict__.copy()


def read_rows(filename: str, content: bytes) -> list[list[str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    text = content.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def guess_mapping(header: list[str]) -> dict[str, int]:
    """Best-effort header → field mapping; the Settings screen lets the owner override it."""
    mapping: dict[str, int] = {}
    lowered = [h.strip().lower() for h in header]
    for fname in FIELDS:
        for i, h in enumerate(lowered):
            if i in mapping.values():
                continue
            if any(s == h or s in h for s in _SYNONYMS[fname]):
                mapping[fname] = i
                break
    return mapping


def _price(value: str) -> float | None:
    cleaned = re.sub(r"[^\d,.\-]", "", value or "")
    if not cleaned:
        return None
    if "," in cleaned and "." not in cleaned:
        cleaned = cleaned.replace(",", ".")
    cleaned = cleaned.replace(",", "")
    try:
        return float(cleaned)
    except ValueError:
        return None


def import_rules(session: Session, filename: str, content: bytes,
                 mapping: dict[str, int] | None = None, is_sample: bool = False) -> tuple[PriceRuleVersion, ImportReport]:
    rows = [r for r in read_rows(filename, content) if any(c.strip() for c in r)]
    if not rows:
        raise ValueError("The spreadsheet is empty")
    header, body = rows[0], rows[1:]
    mapping = mapping or guess_mapping(header)
    for required in ("model", "max_buy_price"):
        if required not in mapping:
            raise ValueError(f"Could not find a '{required}' column — map it manually")

    report, seen, rules = ImportReport(), set(), []
    for row in body:
        cell = lambda f: row[mapping[f]].strip() if f in mapping and mapping[f] < len(row) else ""  # noqa: E731
        raw_model = cell("model")
        model = parse_model(raw_model)
        if not model:
            report.unmatched_models.append(raw_model)
            continue
        storage = parse_storage(cell("storage") + "gb") if cell("storage") else parse_storage(raw_model)
        if storage is None:
            report.missing_storage.append(raw_model)
            continue
        price = _price(cell("max_buy_price"))
        if price is None or price <= 0:
            report.bad_prices.append(f"{model} {storage}GB: '{cell('max_buy_price')}'")
            continue
        key = (model, storage)
        if key in seen:
            report.duplicates.append(f"{model} {storage}GB")
            continue
        seen.add(key)
        battery = _price(cell("min_battery_health"))
        rules.append(PriceRule(model=model, storage_gb=storage, max_buy_price=price,
                               min_battery_health=int(battery) if battery else None,
                               condition_rules=cell("condition_rules") or None,
                               notes=cell("notes") or None))

    report.imported = len(rules)
    for old in session.scalars(select(PriceRuleVersion).where(PriceRuleVersion.is_active)):
        old.is_active = False
    version = PriceRuleVersion(source_name=filename, is_active=True, is_sample=is_sample,
                               report=report.as_dict(), rules=rules)
    session.add(version)
    session.flush()
    return version, report


def import_file(session: Session, path: Path, is_sample: bool = False):
    return import_rules(session, path.name, path.read_bytes(), is_sample=is_sample)


def active_version(session: Session) -> PriceRuleVersion | None:
    return session.scalars(select(PriceRuleVersion).where(PriceRuleVersion.is_active)
                           .order_by(PriceRuleVersion.id.desc())).first()


def lookup(session: Session, model: str | None, storage_gb: int | None) -> PriceRule | None:
    version = active_version(session)
    if not version or not model or not storage_gb:
        return None
    return session.scalars(select(PriceRule).where(
        PriceRule.version_id == version.id, PriceRule.model == model,
        PriceRule.storage_gb == storage_gb)).first()
